#!/usr/bin/env python3
"""
SI-PENA Official Excel Template Generator for OPD Ingestion
BPS Kabupaten Jember

Menghasilkan template berkas Excel (.xlsx) resmi dengan standarisasi format:
1. Standard Table (Tabel Wilayah DDA 31 Kecamatan / KDA Desa) -> Mode DIRECT
2. School Data (Dapodik / EMIS / Kemenag Tingkat Unit) -> Mode AGGREGATE_SCHOOL
3. VKD Questionnaire (Survei Kebutuhan Data 12 Atribut Pelayanan) -> Mode SKD_VKD

Desain visual mengacu pada identitas resmi BPS (Navy Blue #0A3866 & Orange #E67E22).
"""

import sys
import os
import json
import argparse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# 31 Kecamatan Resmi di Kabupaten Jember
JEMBER_DISTRICTS = [
    "Kencong", "Gumukmas", "Puger", "Wuluhan", "Ambulu", "Tempurejo",
    "Silo", "Mayang", "Mumbulsari", "Jenggawah", "Ajung", "Rambipuji",
    "Balung", "Umbulsari", "Semboro", "Jombang", "Sumberbaru", "Tanggul",
    "Bangsalsari", "Panti", "Sukorambi", "Arjasa", "Pakusari", "Kalisat",
    "Ledokombo", "Sumberjambe", "Sukowono", "Jelbuk", "Kaliwates",
    "Sumbersari", "Patrang"
]

# 12 Standar Atribut Pelayanan Publik BPS (VKD Blok III)
SERVICE_ATTRIBUTES = [
    {"code": "U1", "name": "Kesesuaian Persyaratan Layanan", "name_en": "Service Requirements"},
    {"code": "U2", "name": "Kemudahan Prosedur Layanan", "name_en": "Service Procedures"},
    {"code": "U3", "name": "Kecepatan Waktu Pelayanan", "name_en": "Service Time & Speed"},
    {"code": "U4", "name": "Biaya/Tarif Bebas Pungli (Gratis)", "name_en": "Cost/Tariff Zero Fee"},
    {"code": "U5", "name": "Kesesuaian Produk Layanan", "name_en": "Product Specification"},
    {"code": "U6", "name": "Kompetensi Petugas Pelayanan", "name_en": "Officer Competency"},
    {"code": "U7", "name": "Perilaku Sopan & Ramah Petugas", "name_en": "Officer Courtesy & Politeness"},
    {"code": "U8", "name": "Kualitas Sarana & Prasarana PST", "name_en": "Facilities & Infrastructure"},
    {"code": "U9", "name": "Penanganan Pengaduan Konsumen", "name_en": "Complaints Handling"},
    {"code": "U10", "name": "Aksesibilitas Informasi & Website", "name_en": "Data Accessibility & Website"},
    {"code": "U11", "name": "Kemutakhiran Data & Ragam Pilihan", "name_en": "Data Up-to-dateness & Variety"},
    {"code": "U12", "name": "Kenyamanan Ruang Pelayanan", "name_en": "Service Room Comfort"},
]

# Styling Palette BPS
COLOR_NAVY = "0A3866"
COLOR_NAVY_LIGHT = "1B4F72"
COLOR_ORANGE = "E67E22"
COLOR_LIGHT_GRAY = "F8FAFC"
COLOR_BORDER = "CBD5E1"

FONT_TITLE = Font(name="Calibri", size=13, bold=True, color="0A3866")
FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="475569")
FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_HEADER_ORANGE = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_DATA = Font(name="Calibri", size=10, color="1E293B")
FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="1E293B")

FILL_HEADER = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
FILL_HEADER_ORANGE = PatternFill(start_color=COLOR_ORANGE, end_color=COLOR_ORANGE, fill_type="solid")
FILL_ZEBRA = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

THIN_BORDER_SIDE = Side(style="thin", color=COLOR_BORDER)
THIN_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def adjust_column_widths(worksheet, min_width=12):
    for col in worksheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row in [1, 2]:
                continue  # skip title length
            max_len = max(max_len, len(val_str))
        worksheet.column_dimensions[col_letter].width = max(max_len + 4, min_width)


def create_standard_template(output_path, scope="dda", district_name=None, villages=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data_Tabel"
    ws.views.sheetView[0].showGridLines = True

    is_kda = scope.lower() == "kda" or (district_name is not None and district_name != "")
    target_area = f"Kecamatan {district_name}" if (is_kda and district_name) else "Kabupaten Jember"
    level_label = "Desa" if is_kda else "Kecamatan"

    # Header Row (Row 1 for direct parser compatibility)
    headers = [
        "No",
        level_label,
        "Jumlah_Unit",
        "Kapasitas_Volume",
        "Realisasi_Capaian",
        "Keterangan"
    ]

    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h_text)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 28

    # Populate Rows
    if is_kda and villages:
        row_items = [v.strip() for v in villages if v.strip()]
    elif is_kda and not villages:
        row_items = [f"Desa {i+1}" for i in range(8)]
    else:
        row_items = JEMBER_DISTRICTS

    current_row = 2
    for idx, item in enumerate(row_items, start=1):
        c1 = ws.cell(row=current_row, column=1, value=idx)
        c2 = ws.cell(row=current_row, column=2, value=item)
        c3 = ws.cell(row=current_row, column=3, value="")
        c4 = ws.cell(row=current_row, column=4, value="")
        c5 = ws.cell(row=current_row, column=5, value="")
        c6 = ws.cell(row=current_row, column=6, value="")

        for c in [c1, c2, c3, c4, c5, c6]:
            c.font = FONT_DATA
            c.border = THIN_BORDER

        c1.alignment = ALIGN_CENTER
        c2.alignment = ALIGN_LEFT
        c3.alignment = ALIGN_RIGHT
        c4.alignment = ALIGN_RIGHT
        c5.alignment = ALIGN_RIGHT
        c6.alignment = ALIGN_LEFT

        if idx % 2 == 0:
            for c in [c1, c2, c3, c4, c5, c6]:
                c.fill = FILL_ZEBRA

        current_row += 1

    # Total Row
    tot_row = current_row
    ws.cell(row=tot_row, column=1, value="").border = THIN_BORDER
    c_tot_label = ws.cell(row=tot_row, column=2, value="JUMLAH / TOTAL")
    c_tot_label.font = FONT_BOLD
    c_tot_label.alignment = ALIGN_LEFT
    c_tot_label.border = THIN_BORDER

    for col_idx in [3, 4, 5]:
        col_letter = get_column_letter(col_idx)
        c_tot = ws.cell(row=tot_row, column=col_idx, value=f"=SUM({col_letter}2:{col_letter}{tot_row-1})")
        c_tot.font = FONT_BOLD
        c_tot.alignment = ALIGN_RIGHT
        c_tot.border = THIN_BORDER

    ws.cell(row=tot_row, column=6, value="").border = THIN_BORDER

    adjust_column_widths(ws)

    # Instruction Sheet
    ws_guide = wb.create_sheet(title="Petunjuk_Pengisian")
    ws_guide.views.sheetView[0].showGridLines = True
    ws_guide.cell(row=1, column=1, value=f"PANDUAN PENGISIAN DATA TABEL ({target_area.upper()})").font = FONT_TITLE
    ws_guide.cell(row=2, column=1, value="Standar Ingesti Berkas Mentah OPD - SI-PENA BPS Kabupaten Jember").font = FONT_SUBTITLE

    instructions = [
        "1. Jangan mengubah nama kolom pada baris header (Baris ke-1 pada sheet Data_Tabel).",
        "2. Masukkan angka secara langsung tanpa menyisipkan satuan teks seperti 'orang', 'ton', atau 'unit' di dalam sel angka.",
        "3. Bila terdapat data desimal, Anda dapat menggunakan tanda koma (,) atau titik (.). Engine SI-PENA akan menormalkannya secara otomatis.",
        "4. Bila data kosong atau tidak ada kegiatan, masukkan tanda strip (-) atau angka 0.",
        "5. DILARANG menggabungkan sel (merge cells) di dalam area data angka karena dapat mengganggu proses ekstraksi Python.",
        "6. Anda diperbolehkan mengubah nama kolom 'Jumlah_Unit', 'Kapasitas_Volume', dll. sesuai indikator riil OPD Anda.",
        "7. Simpan berkas dalam format .xlsx atau .xls sebelum diunggah ke portal SI-PENA."
    ]

    for r_idx, text in enumerate(instructions, start=4):
        c = ws_guide.cell(row=r_idx, column=1, value=text)
        c.font = FONT_DATA
    ws_guide.column_dimensions["A"].width = 95

    wb.save(output_path)
    return output_path


def create_schools_template(output_path, district_name=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data_Sekolah"
    ws.views.sheetView[0].showGridLines = True

    # Header Row (Row 1 for direct parser compatibility)
    headers = [
        "No",
        "Nama_Sekolah",
        "Kecamatan",
        "Desa",
        "Jenjang",
        "Status",
        "Guru",
        "Murid"
    ]

    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h_text)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 28

    # Sample rows for guidance
    default_dist = district_name or "Ambulu"
    samples = [
        (1, f"SDN {default_dist} 01", default_dist, "Karanganyar", "SD", "Negeri", 14, 285),
        (2, f"SDN {default_dist} 02", default_dist, "Sabrang", "SD", "Negeri", 12, 240),
        (3, f"SD Swasta Harapan", default_dist, "Tegalsari", "SD", "Swasta", 10, 180),
        (4, f"SMPN 1 {default_dist}", default_dist, "Pontang", "SMP", "Negeri", 32, 620),
        (5, f"SMP Swasta Islam", default_dist, "Andongsari", "SMP", "Swasta", 18, 290),
        (6, f"SMAN 1 {default_dist}", default_dist, "Sumberrejo", "SMA", "Negeri", 45, 840),
        (7, f"SMK Teknologi {default_dist}", default_dist, "Sabrang", "SMK", "Swasta", 28, 410),
    ]

    for row_idx, s in enumerate(samples, start=2):
        for col_idx, val in enumerate(s, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = FONT_DATA
            cell.border = THIN_BORDER
            if col_idx in [1, 5, 6]:
                cell.alignment = ALIGN_CENTER
            elif col_idx in [7, 8]:
                cell.alignment = ALIGN_RIGHT
            else:
                cell.alignment = ALIGN_LEFT

            if row_idx % 2 == 1:
                cell.fill = FILL_ZEBRA

    adjust_column_widths(ws)

    # Add Dropdown Validations for Jenjang & Status
    dv_jenjang = DataValidation(type="list", formula1='"SD,SMP,SMA,SMK,MI,MTs,MA,SLB"', allow_blank=True)
    ws.add_data_validation(dv_jenjang)
    dv_jenjang.add("E2:E300")

    dv_status = DataValidation(type="list", formula1='"Negeri,Swasta"', allow_blank=True)
    ws.add_data_validation(dv_status)
    dv_status.add("F2:F300")

    # Instruction Sheet
    ws_guide = wb.create_sheet(title="Petunjuk_Agregasi")
    ws_guide.views.sheetView[0].showGridLines = True
    ws_guide.cell(row=1, column=1, value="PANDUAN AGREGASI DATA SEKOLAH (DAPODIK / EMIS)").font = FONT_TITLE
    ws_guide.cell(row=2, column=1, value="Standar Ingesti Agregasi Pendidikan SI-PENA BPS Kabupaten Jember").font = FONT_SUBTITLE

    instructions = [
        "1. Berkas ini digunakan untuk mengagregasi data sekolah individu ke rekapitulasi Bab 4 (Sosial).",
        "2. Kolom 'Kecamatan' wajib diisi untuk agregasi level Kabupaten (DDA).",
        "3. Kolom 'Desa' wajib diisi untuk agregasi level Kecamatan (KDA).",
        "4. Pilihan 'Jenjang': SD, SMP, SMA, SMK, MI, MTs, MA, SLB.",
        "5. Pilihan 'Status': Negeri atau Swasta.",
        "6. Kolom 'Guru' dan 'Murid' harus berupa angka bulat (integer).",
        "7. Engine Python SI-PENA akan otomatis menghitung jumlah unit sekolah, total guru, total murid, dan breakdown per status."
    ]

    for r_idx, text in enumerate(instructions, start=4):
        c = ws_guide.cell(row=r_idx, column=1, value=text)
        c.font = FONT_DATA
    ws_guide.column_dimensions["A"].width = 95

    wb.save(output_path)
    return output_path


def create_skd_template(output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kuesioner_VKD"
    ws.views.sheetView[0].showGridLines = True

    # Header Row (Row 1 for direct parser compatibility)
    # Column A: No_Responden
    c_resp = ws.cell(row=1, column=1, value="No_Responden")
    c_resp.font = FONT_HEADER
    c_resp.fill = FILL_HEADER
    c_resp.alignment = ALIGN_CENTER
    c_resp.border = THIN_BORDER

    col_idx = 2
    for attr in SERVICE_ATTRIBUTES:
        # X: Kepuasan (Navy)
        cX = ws.cell(row=1, column=col_idx, value=f"{attr['code']}_X")
        cX.font = FONT_HEADER
        cX.fill = FILL_HEADER
        cX.alignment = ALIGN_CENTER
        cX.border = THIN_BORDER

        # Y: Kepentingan (Orange)
        cY = ws.cell(row=1, column=col_idx + 1, value=f"{attr['code']}_Y")
        cY.font = FONT_HEADER_ORANGE
        cY.fill = FILL_HEADER_ORANGE
        cY.alignment = ALIGN_CENTER
        cY.border = THIN_BORDER

        col_idx += 2

    # Saran / Catatan column
    cSaran = ws.cell(row=1, column=col_idx, value="Saran_Masukan")
    cSaran.font = FONT_HEADER
    cSaran.fill = FILL_HEADER
    cSaran.alignment = ALIGN_CENTER
    cSaran.border = THIN_BORDER

    ws.row_dimensions[1].height = 28

    # Populate 10 sample respondents
    sample_scores = [
        [4, 4, 3, 4, 4, 4, 4, 4, 3, 4, 4, 4, 4, 4, 3, 4, 4, 4, 4, 4, 3, 4, 4, 4],
        [3, 4, 4, 4, 4, 4, 4, 4, 3, 3, 4, 4, 4, 4, 4, 4, 3, 3, 4, 4, 4, 4, 3, 4],
        [4, 4, 4, 4, 3, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 3, 4, 4, 4, 4, 4],
        [3, 3, 3, 4, 4, 4, 4, 4, 3, 4, 3, 4, 4, 4, 3, 4, 3, 4, 4, 4, 3, 4, 4, 4],
        [4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4],
        [3, 4, 3, 4, 3, 4, 4, 4, 3, 4, 3, 4, 3, 4, 3, 4, 3, 4, 3, 4, 3, 4, 3, 4],
        [4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4],
        [4, 4, 3, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 3, 4, 4, 4, 4, 4, 4, 4],
        [3, 4, 4, 4, 3, 4, 4, 4, 3, 4, 4, 4, 4, 4, 3, 4, 3, 4, 4, 4, 3, 4, 4, 4],
        [4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4, 4],
    ]

    for resp_idx, scores in enumerate(sample_scores, start=1):
        row_num = resp_idx + 1
        c_r = ws.cell(row=row_num, column=1, value=resp_idx)
        c_r.font = FONT_DATA
        c_r.alignment = ALIGN_CENTER
        c_r.border = THIN_BORDER

        for s_idx, val in enumerate(scores, start=2):
            cell = ws.cell(row=row_num, column=s_idx, value=val)
            cell.font = FONT_DATA
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER

        c_saran = ws.cell(row=row_num, column=len(scores) + 2, value="Pelayanan ramah dan cepat." if resp_idx == 1 else "")
        c_saran.font = FONT_DATA
        c_saran.alignment = ALIGN_LEFT
        c_saran.border = THIN_BORDER

        if resp_idx % 2 == 0:
            c_r.fill = FILL_ZEBRA
            for s_idx in range(2, len(scores) + 3):
                ws.cell(row=row_num, column=s_idx).fill = FILL_ZEBRA

    adjust_column_widths(ws, min_width=8)

    # Sheet 2: Guide & Reference for 12 Attributes
    ws_guide = wb.create_sheet(title="Panduan_12_Atribut")
    ws_guide.views.sheetView[0].showGridLines = True
    ws_guide.cell(row=1, column=1, value="DAFTAR 12 ATRIBUT PELAYANAN PUBLIK BPS (VKD BLOK III)").font = FONT_TITLE
    ws_guide.cell(row=2, column=1, value="Skala Penilaian: 1 = Tidak Puas/Penting, 2 = Kurang, 3 = Puas/Penting, 4 = Sangat Puas/Penting").font = FONT_SUBTITLE

    headers_guide = ["Kode", "Nama Indikator Pelayanan", "Keterangan Variabel X & Y"]
    for col_idx, h_text in enumerate(headers_guide, start=1):
        cell = ws_guide.cell(row=4, column=col_idx, value=h_text)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER
    ws_guide.row_dimensions[4].height = 24

    for r_idx, attr in enumerate(SERVICE_ATTRIBUTES, start=5):
        c1 = ws_guide.cell(row=r_idx, column=1, value=attr["code"])
        c2 = ws_guide.cell(row=r_idx, column=2, value=attr["name"])
        c3 = ws_guide.cell(row=r_idx, column=3, value=f"{attr['code']}_X (Tingkat Kepuasan) & {attr['code']}_Y (Tingkat Kepentingan)")

        for c in [c1, c2, c3]:
            c.font = FONT_DATA
            c.border = THIN_BORDER

        c1.alignment = ALIGN_CENTER
        c2.alignment = ALIGN_LEFT
        c3.alignment = ALIGN_LEFT

        if (r_idx - 4) % 2 == 0:
            for c in [c1, c2, c3]:
                c.fill = FILL_ZEBRA

    ws_guide.column_dimensions["A"].width = 12
    ws_guide.column_dimensions["B"].width = 45
    ws_guide.column_dimensions["C"].width = 50

    wb.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="SI-PENA Official Excel Template Generator")
    parser.add_argument("--type", choices=["standard", "schools", "skd"], required=True, help="Tipe template Excel")
    parser.add_argument("--output", required=True, help="Path target penyimpanan file Excel (.xlsx)")
    parser.add_argument("--scope", choices=["dda", "kda"], default="dda", help="Cakupan wilayah: dda (kabupaten) atau kda (kecamatan)")
    parser.add_argument("--district", help="Nama kecamatan (khusus KDA)")
    parser.add_argument("--villages", help="Daftar desa dipisah koma (khusus KDA)")
    args = parser.parse_args()

    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)

    try:
        if args.type == "standard":
            villages_list = [v.strip() for v in args.villages.split(",")] if args.villages else None
            create_standard_template(args.output, scope=args.scope, district_name=args.district, villages=villages_list)
        elif args.type == "schools":
            create_schools_template(args.output, district_name=args.district)
        elif args.type == "skd":
            create_skd_template(args.output)

        print(json.dumps({
            "status": "success",
            "type": args.type,
            "output_path": args.output,
            "file_size": os.path.getsize(args.output)
        }, ensure_ascii=False, indent=2))

    except Exception as e:
        print(json.dumps({
            "status": "error",
            "message": str(e)
        }, ensure_ascii=False, indent=2))
        sys.exit(1)


if __name__ == "__main__":
    main()
