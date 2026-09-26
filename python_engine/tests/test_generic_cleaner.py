"""Pengujian pembersih Excel OPD (sheet resolution, merge, normalisasi)."""
import json
import subprocess
import sys

import openpyxl
import pytest

import generic_cleaner as gc

PY = sys.executable
CLEANER = str(__import__("pathlib").Path(gc.__file__).resolve())


def make_workbook(path, sheets):
    """sheets: dict nama_sheet -> list baris (list sel)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    wb.save(path)
    return path


def test_resolve_sheet_by_index(tmp_path):
    wb = openpyxl.load_workbook(make_workbook(tmp_path / "a.xlsx", {
        "Depan": [["A", "B"]],
        "Belakang": [["C", "D"]],
    }))
    assert gc.resolve_sheet(wb, "0").title == "Depan"
    assert gc.resolve_sheet(wb, "1").title == "Belakang"
    assert gc.resolve_sheet(wb, None).title == "Depan"


def test_resolve_sheet_by_name_and_error_lists_sheets(tmp_path):
    wb = openpyxl.load_workbook(make_workbook(tmp_path / "b.xlsx", {
        "Rekap": [["A"]],
        "Mentah": [["B"]],
    }))
    assert gc.resolve_sheet(wb, "Mentah").title == "Mentah"

    with pytest.raises(ValueError) as err:
        gc.resolve_sheet(wb, "TidakAda")
    assert "Rekap" in str(err.value) and "Mentah" in str(err.value)

    with pytest.raises(ValueError):
        gc.resolve_sheet(wb, "9")


def test_clean_excel_merges_fills_and_resets_index(tmp_path):
    path = make_workbook(tmp_path / "c.xlsx", {"Sheet1": [
        ["Kecamatan", "Desa", "Guru", "Murid"],
        ["Kencong", "Cakru", "10", "200"],
        [None, "Sukorejo", "12", "250"],   # A3 digabung dengan A2 -> forward fill
        [None, None, None, None],           # baris kosong harus dibuang
        ["Kencong", "Pengatang", "8", "180"],
    ]})
    wb = openpyxl.load_workbook(path)
    wb["Sheet1"].merge_cells("A2:A3")
    wb.save(path)

    result = gc.clean_excel(path, sheet_name="0")

    assert result["status"] == "success"
    assert result["headers"] == ["Kecamatan", "Desa", "Guru", "Murid"]
    # reset_index dipakai: baris data tidak boleh hilang/ganjil karena index lama
    assert result["total_rows"] == 3
    assert result["data"][0]["Kecamatan"] == "Kencong"
    assert result["data"][0]["Murid"] == 200.0
    # nilai merge ter-forward-fill, bukan kosong
    assert result["data"][1]["Kecamatan"] == "Kencong"
    assert result["data"][1]["Desa"] == "Sukorejo"
    assert result["data"][2]["Murid"] == 180.0


def test_clean_excel_normalizes_decimal_comma(tmp_path):
    path = make_workbook(tmp_path / "d.xlsx", {"Sheet1": [
        ["Komoditas", "Luas", "Produksi"],
        ["Padi", "1.250,5", "48.210"],
    ]})
    result = gc.clean_excel(path, sheet_name="0")
    assert result["status"] == "success"
    row = result["data"][0]
    assert row["Luas"] == pytest.approx(1250.5)
    assert row["Produksi"] == pytest.approx(48210.0)


def test_clean_excel_unknown_sheet_returns_error(tmp_path):
    path = make_workbook(tmp_path / "e.xlsx", {"Data": [["A"]]})
    result = gc.clean_excel(path, sheet_name="Hilang")
    assert result["status"] == "error"
    assert "Data" in result["message"]


def test_cli_emits_json(tmp_path):
    path = make_workbook(tmp_path / "f.xlsx", {"Sheet1": [["A", "B"], ["1", "2"]]})
    proc = subprocess.run(
        [PY, CLEANER, str(path), "--sheet", "Sheet1"],
        capture_output=True, text=True, encoding="utf-8", timeout=60,
    )
    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    assert payload["status"] == "success"
    assert payload["sheet_name"] == "Sheet1"
