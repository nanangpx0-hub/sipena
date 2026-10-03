import os
import openpyxl
from python_engine.generators.excel_templates import (
    create_standard_template,
    create_schools_template,
    create_skd_template,
    JEMBER_DISTRICTS,
    SERVICE_ATTRIBUTES,
)
from python_engine.parsers.generic_cleaner import clean_excel
from python_engine.parsers.individual_aggregator import aggregate_school_data
from python_engine.calculators.skd_engine import calculate_skd_metrics
import pandas as pd


def test_standard_template_generation_and_cleaning(tmp_path):
    out_file = str(tmp_path / "standard_test.xlsx")
    create_standard_template(out_file, scope="dda")

    assert os.path.exists(out_file)
    wb = openpyxl.load_workbook(out_file)
    assert "Data_Tabel" in wb.sheetnames
    assert "Petunjuk_Pengisian" in wb.sheetnames

    # Clean with generic cleaner
    result = clean_excel(out_file)
    assert result["status"] == "success"
    assert "Kecamatan" in result["headers"]
    # Jember districts + total row
    assert len(result["data"]) >= len(JEMBER_DISTRICTS)


def test_standard_template_kda_with_villages(tmp_path):
    out_file = str(tmp_path / "kda_ambulu.xlsx")
    villages = ["Andongsari", "Karanganyar", "Pontang", "Sabrang", "Sumberrejo", "Tegalsari", "Ambulu"]
    create_standard_template(out_file, scope="kda", district_name="Ambulu", villages=villages)

    assert os.path.exists(out_file)
    result = clean_excel(out_file)
    assert result["status"] == "success"
    assert "Desa" in result["headers"]
    first_item = result["data"][0]["Desa"]
    assert first_item == "Andongsari"


def test_schools_template_and_aggregation(tmp_path):
    out_file = str(tmp_path / "schools_test.xlsx")
    create_schools_template(out_file, district_name="Kencong")

    assert os.path.exists(out_file)
    res = aggregate_school_data(out_file, level="kecamatan")
    assert res["status"] == "success"
    assert res["summary"]["total_sekolah"] == 7
    assert res["summary"]["total_guru"] > 0
    assert res["summary"]["total_murid"] > 0


def test_skd_template_and_engine(tmp_path):
    out_file = str(tmp_path / "skd_test.xlsx")
    svg_file = str(tmp_path / "cartesian.svg")
    create_skd_template(out_file)

    assert os.path.exists(out_file)
    df = pd.read_excel(out_file)
    assert "U1_X" in df.columns
    assert "U12_Y" in df.columns

    res = calculate_skd_metrics(df, output_svg_path=svg_file)
    assert res["status"] == "success"
    assert "ikk_score" in res
    assert "ipak_score" in res
    assert os.path.exists(svg_file)
