#!/usr/bin/env python3
"""
SI-PENA Generic Excel Cleaner
Pembersih data mentah Excel OPD:
- Unmerge cells dan foward-fill nilai merge
- Menghapus baris dan kolom yang kosong
- Normalisasi pemisah desimal koma (,) ke titik (.)
- Ekspor tabel bersih terstruktur ke JSON stdout
"""

import sys
import json
import argparse
import pandas as pd
import openpyxl

def unmerge_and_fill(workbook):
    for sheet in workbook.worksheets:
        merged_ranges = list(sheet.merged_cells.ranges)
        for cell_range in merged_ranges:
            min_col, min_row, max_col, max_row = cell_range.bounds
            top_left_value = sheet.cell(row=min_row, column=min_col).value
            sheet.unmerge_cells(start_row=min_row, start_column=min_col,
                                end_row=max_row, end_column=max_col)
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    sheet.cell(row=row, column=col).value = top_left_value
    return workbook

def resolve_sheet(workbook, sheet_name):
    """Resolve argumen --sheet menjadi worksheet (mendukung indeks '0' maupun nama sheet)."""
    available = [ws.title for ws in workbook.worksheets]
    if sheet_name is None:
        return workbook.worksheets[0]

    token = str(sheet_name).strip()
    if token.lstrip('-').isdigit():
        idx = int(token)
        if idx < 0 or idx >= len(workbook.worksheets):
            raise ValueError(
                "Indeks sheet %d di luar jangkauan (tersedia: %s)" % (idx, ", ".join(available))
            )
        return workbook.worksheets[idx]

    if token in workbook.sheetnames:
        return workbook[token]

    raise ValueError(
        "Sheet '%s' tidak ditemukan (tersedia: %s)" % (token, ", ".join(available))
    )


def clean_excel(file_path, sheet_name=0, header_row=None):
    try:
        # Load workbook with openpyxl to resolve merge cells
        wb = openpyxl.load_workbook(file_path, data_only=True)
        wb = unmerge_and_fill(wb)

        # Pilih sheet: terima indeks numerik ("0") maupun nama sheet
        active_sheet = resolve_sheet(wb, sheet_name)
        data = list(active_sheet.values)

        if not data:
            return {"status": "error", "message": "Worksheet is empty"}

        df = pd.DataFrame(data)

        # Drop fully empty rows and columns
        df.dropna(how='all', inplace=True)
        df.dropna(axis=1, how='all', inplace=True)
        df.reset_index(drop=True, inplace=True)

        # Detect header row: row with maximum non-null string values
        if header_row is None:
            best_idx = 0
            best_score = -1
            for idx, row in df.head(10).iterrows():
                score = sum(1 for val in row if isinstance(val, str) and len(str(val).strip()) > 0)
                if score > best_score:
                    best_score = score
                    best_idx = idx
            header_row = best_idx

        # Assign header
        header = [str(col).strip() if col is not None else f"Kolom_{i+1}" for i, col in enumerate(df.loc[header_row])]
        df = df.iloc[header_row + 1:].copy()
        df.columns = header

        # Drop rows where all columns are empty
        df.dropna(how='all', inplace=True)

        # Normalize decimals and numbers
        cleaned_records = []
        for _, row in df.iterrows():
            record = {}
            for col, val in row.items():
                if val is None:
                    record[col] = "-"
                elif isinstance(val, (int, float)):
                    record[col] = float(val) if isinstance(val, float) else int(val)
                else:
                    str_val = str(val).strip()
                    # Try replacing comma with period for decimal numbers
                    if str_val.replace(',', '').replace('.', '').replace('-', '').isdigit():
                        try:
                            num = float(str_val.replace('.', '').replace(',', '.'))
                            record[col] = num
                        except ValueError:
                            record[col] = str_val
                    else:
                        record[col] = str_val
            cleaned_records.append(record)

        return {
            "status": "success",
            "sheet_name": active_sheet.title,
            "headers": header,
            "total_rows": len(cleaned_records),
            "data": cleaned_records
        }

    except Exception as e:
        return {"status": "error", "message": str(e)}

def main():
    parser = argparse.ArgumentParser(description="SI-PENA Generic Excel Cleaner")
    parser.add_argument("file_path", help="Path ke berkas Excel mentah OPD")
    parser.add_argument("--sheet", default=0, help="Nama atau indeks sheet (default 0)")
    parser.add_argument("--header", type=int, default=None, help="Indeks baris header")
    args = parser.parse_args()

    result = clean_excel(args.file_path, args.sheet, args.header)
    print(json.dumps(result, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
